# -*- coding: utf-8 -*-
"""거제·통영·사천·김해(바로물류) 게시판 데이터: 쿠팡 어드민 전체 다운로드 + 카톡 배송 스케줄 표 → region-data.js

사용법:
  python tools/convert_region.py

- 전체스케쥴(어드민): source/region/admin/*.xlsx — 쿠팡 어드민 '엑셀 다운로드' 전체본(모든 캠프가 한 파일에 들어 있음).
  같은 날짜가 여러 파일에 있으면 파일명 끝 숫자(내보낸 시각)가 나중인 파일이 그 날짜를 통째로 대체(convert_day.py 와 같음).
  이름은 어드민 그대로(아이디 주인 이름).
- 배송 스케줄: source/region/plan.json 에 적은 카톡 표(지역·주간/야간·파일·시트). 파일은 source/region/plan/ 에.
  표 모양 두 가지를 읽음:
    ① 7일 가로 블록 — 'No.' 머리행, 블록 10칸(No.|주/야|표시|이름|ID|라우트|서브라우트|전화|차량).
       블록 위 제목에서 날짜, 제목에 'V_통영2'면 V_통영2, '거제'면 M_거제1. 블록 안 'V_통영0 …' 소제목 아래는 V_통영0.
    ② 일일입차예정(야간) — 'MM월DD일 X요일' 날짜 칸, 하루 4칸(표시|이름|ID|라우트). 캠프는 지역 기준(통영 M_통영1, 거제 M_거제1).
  표시 칸: '' · '백업' · 그 밖의 글자 = 출근, '휴무'(또는 라우트 칸이 '휴무') = 휴무.
  이름은 표에 적힌 그대로(실제 운행자). 라우트는 '라우트' 칸만(서브라우트 무시), 읽을 수 없는 표기는 글자 그대로 보여줌.
  같은 날·캠프·아이디가 여러 표에 있으면 plan.json 뒤쪽 것이 이김(내용이 다르면 알림).
  plan.json = {"files": [{region, group, file, sheet}, ...], "name_fix": {"표 오타": "맞는 이름"}} — name_fix 는 본부장님이 확인한 것만.
  표시 칸에 '휴무'·'백업'·빈칸이 아닌 글자(예: 다른 사람 이름)가 있으면 출근으로 두고 알림만 띄움.
- 동명이인(같은 이름에 아이디가 여럿): 이름 뒤에 그 사람의 주 라우트를 붙임 — 예) 김상우(607A), 김상우(608A).
- 아이디는 같은 사람을 가리는 데만 쓰고 내보내지 않음. 전화번호·차량번호·사업자번호도 내보내지 않음 (repo 공개).
  원본은 git 제외 폴더 source/ 에만.
- 결과: region-data.js → window.REGION_DATA = {generated, regions: [{key, name, day: {plan, admin}, night: {plan, admin}}]}
  plan/admin 은 index.html makeFull() 이 그리는 모양({title, dates, weeks}) 또는 null.

■ 매주 절차
  1. 본부장님이 주는 어드민 전체 다운로드(schedule-v2-…_<숫자>.xlsx)를 source/region/admin/ 에 넣음 (이름 그대로).
  2. 카톡 배송 스케줄 표는 source/region/plan/ 에 넣고 plan.json 에 {region, group, file, sheet} 한 줄 추가.
  3. python tools/convert_region.py → region-data.js 갱신 → git add region-data.js && commit && push
"""
import sys, os, io, re, json, glob, datetime, collections, warnings
import openpyxl

warnings.filterwarnings('ignore')
sys.stdout.reconfigure(encoding='utf-8')
ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
sys.path.insert(0, os.path.join(ROOT, 'tools'))
import convert_day as cd                                    # 주별 묶기(build_week)·파일 순서(stamp)는 부산 주간과 같은 코드

SRC = os.path.join(ROOT, 'source', 'region')
OUT = os.environ.get('REGION_OUT') or os.path.join(ROOT, 'region-data.js')
REGIONS = [   # 게시판 순서. 캠프는 화면에 이 순서로
    {'key': 'geoje', 'name': '거제', 'day': ['M_거제1'], 'night': ['M_거제1']},
    {'key': 'tongyeong', 'name': '통영', 'day': ['V_통영2', 'V_통영0'], 'night': ['M_통영1']},
    {'key': 'sacheon', 'name': '사천', 'day': ['M_사천1'], 'night': ['M_사천1']},
    {'key': 'gimhae', 'name': '김해', 'day': ['김해1', '김해2'], 'night': ['김해1', '김해2']},   # 바로물류 (10/26~)
]
WAVES = {'day': ('WAVE2', {'', '1회전,2회전'}), 'night': ('WAVE1', {'', 'D1,D2,F3'})}
ROUTE_FIX = {'0322C': '032C'}                               # 통영 표 오타
ROUTE_RE = re.compile(r'\d{3}[A-Z]')
DR = re.compile(r'(\d{4})년\s*(\d{1,2})월\s*(\d{1,2})일')
DT = re.compile(r'^\s*(\d{1,2})월\s*(\d{1,2})일\s*([월화수목금토일])요일\s*$')
WD = '월화수목금토일'
# 한 줄 = (날짜, 캠프, 이름, 상태, 회전, [라우트], 아이디)  — 아이디는 내부에서만 씀


def fail(msg, items=()):
    print(f'[중단] {msg} — {os.path.basename(OUT)}를 쓰지 않았습니다.')
    for x in list(items)[:30]:
        print('   ✗', x)
    sys.exit(1)


def routes_of(text):
    out = []
    for t in re.split(r'[,.\s]+', str(text or '').strip().upper()):
        if t:
            t = ROUTE_FIX.get(t, t)
            out += re.findall(r'\d{3}[A-Z]', t) if re.fullmatch(r'(?:\d{3}[A-Z]){2,}', t) else [t]   # 쉼표 빠진 '704D705B' → 704D, 705B
    return out


# ---------------- 어드민 (아이디까지 읽음) ----------------
def load_admin(path, camps, wave):
    ws = openpyxl.load_workbook(path, read_only=True, data_only=True).worksheets[0]
    it = ws.iter_rows(values_only=True)
    hdr = [str(c).strip() if c is not None else '' for c in next(it)]
    ix = {h: i for i, h in enumerate(hdr)}
    need = ['업무일', '캠프명', '웨이브', '이름', '아이디', '업무상태', '회전', '업무라우트']
    miss = [h for h in need if h not in ix]
    if miss:
        fail(f"'{os.path.basename(path)}' 어드민 양식 머리행이 아님 (없는 칸: {miss})")
    rows, errs = [], []
    for n, r in enumerate(it, start=2):
        if not any(v not in (None, '') for v in r):
            continue
        g = lambda k: r[ix[k]] if ix[k] < len(r) else None
        camp, wv = str(g('캠프명') or '').strip(), str(g('웨이브') or '').strip().upper()
        if camp not in camps or not wv.startswith(wave):
            continue
        where = f'{os.path.basename(path)} {n}행'
        d = cd.as_date(g('업무일'), where)
        name, pid = str(g('이름') or '').strip(), str(g('아이디') or '').strip()
        st = str(g('업무상태') or '').strip()
        rot = re.sub(r'\s+', '', str(g('회전') or ''))
        routes = [x for x in re.split(r'[,\s]+', str(g('업무라우트') or '').strip().upper()) if x]
        if not name:
            errs.append(f'{where}: 이름 빈칸')
        elif st == '출근':
            if not routes or any(not ROUTE_RE.fullmatch(x) for x in routes):
                errs.append(f"{where} {d} {name}: 출근인데 업무라우트 '{g('업무라우트')}' 이상")
        elif st == '휴무':
            if routes:
                errs.append(f"{where} {d} {name}: 휴무인데 업무라우트 있음")
        else:
            errs.append(f"{where} {d} {name}: 업무상태 '{st}' (출근/휴무 만)")
        rows.append((d, camp, name, st, rot, routes, pid or name))
    if errs:
        fail(f"'{os.path.basename(path)}' 검증 실패 {len(errs)}건", errs)
    return rows


def admin_rows(camps, wave):
    """→ (줄, {날짜: 그 날짜를 이긴 다운로드의 내보낸 시각(초)})"""
    by_date, when = {}, {}
    for p in sorted(glob.glob(os.path.join(SRC, 'admin', '*.xlsx')), key=cd.stamp):   # 오래된 → 최신
        if os.path.basename(p).startswith('~$'):
            continue
        rows = load_admin(p, camps, wave)
        for d in sorted({x[0] for x in rows}):
            by_date[d] = [x for x in rows if x[0] == d]
            when[d] = cd.stamp(p) / 1000
    return [x for d in sorted(by_date) for x in by_date[d]], when


# ---------------- 카톡 표 ① 7일 가로 블록 ----------------
def parse_blocks(ws, region, group):
    hdr = next((r for r in range(1, 15) if any(str(ws.cell(r, c).value or '').strip() == 'No.' for c in range(1, ws.max_column + 1))), None)
    if hdr is None:
        return None
    starts = [c for c in range(1, ws.max_column + 1) if str(ws.cell(hdr, c).value or '').strip() == 'No.']
    rows = []
    for c0 in starts:
        title = ' '.join(str(ws.cell(r, c).value) for r in range(1, hdr) for c in range(c0, c0 + 10) if ws.cell(r, c).value)
        m = DR.search(title)
        if not m:
            fail(f'{ws.title}: {openpyxl.utils.get_column_letter(c0)}열 블록 제목에 날짜가 없음')
        d = datetime.date(*map(int, m.groups())).isoformat()
        camp = 'V_통영2' if 'V_통영2' in title else 'M_거제1' if '거제' in title else region[group][0]
        for r in range(hdr + 1, ws.max_row + 1):
            no = ws.cell(r, c0).value
            if isinstance(no, str) and 'V_통영0' in no:
                camp = 'V_통영0'
                continue
            nm = ws.cell(r, c0 + 3).value
            if not isinstance(no, (int, float)) or not nm:
                continue
            nm = str(nm).strip()
            pid = str(ws.cell(r, c0 + 4).value or '').strip() or nm
            mark = str(ws.cell(r, c0 + 2).value or '').strip()
            rt = ws.cell(r, c0 + 5).value
            if mark not in ('', '휴무', '백업'):
                print(f'     (참고) {ws.title} {d} {nm}: 표시 칸 "{mark}" — 출근으로 둠')
            off = mark == '휴무' or str(rt or '').strip() == '휴무'
            rows.append((d, camp, nm, '휴무' if off else '출근', '', [] if off else routes_of(rt), pid))
    return rows


# ---------------- 카톡 표 ② 일일입차예정(야간, 하루 4칸) ----------------
def parse_daily4(ws, region, group, year):
    heads = []
    for r in range(1, ws.max_row + 1):
        for c in range(1, ws.max_column + 1):
            v = ws.cell(r, c).value
            m = DT.match(v) if isinstance(v, str) else None
            if m:
                d = datetime.date(year, int(m.group(1)), int(m.group(2)))
                heads.append((r, c, d if WD[d.weekday()] == m.group(3) else None))
    if not heads:
        return None
    hrows = sorted({r for r, c, d in heads})
    camp = region[group][0]
    rows = []
    for r0, c0, d in heads:
        nxt = [x for x in hrows if x > r0]
        r_end = (nxt[0] - 1) if nxt else ws.max_row
        block = []
        for r in range(r0 + 1, r_end + 1):
            mark, nm, pid, rt = (ws.cell(r, c0 + k).value for k in range(4))
            if not nm or not isinstance(nm, str) or nm.replace(' ', '') == '백업':
                continue
            off = str(mark or '').strip() == '휴무' or str(rt or '').strip() == '휴무'
            if str(mark or '').strip() not in ('', '휴무', '백업'):
                print(f'     (참고) {ws.title} {d} {nm.strip()}: 표시 칸 "{str(mark).strip()}" — 출근으로 둠')
            block.append((nm.strip(), off, rt, str(pid or '').strip() or nm.strip()))
        if not block:
            continue                                       # 빈 블록(템플릿에 남은 '6월4일 목요일' 등)
        if d is None:
            fail(f'{ws.title}: 요일이 안 맞는 날짜 블록에 내용이 있음')
        for nm, off, rt, pid in block:
            rows.append((d.isoformat(), camp, nm, '휴무' if off else '출근', '', [] if off else routes_of(rt), pid))
    return rows


def plan_rows(region, group, items, name_fix):
    """→ (줄, {날짜: 그 날짜 줄을 준 카톡 표 파일을 source 에 넣은 시각(초) 중 가장 늦은 것})"""
    merged, notes, src_t = {}, [], {}
    for it in items:
        path = os.path.join(SRC, 'plan', it['file'])
        if not os.path.exists(path):
            fail(f"plan.json 의 파일이 없음: {it['file']}")
        wb = openpyxl.load_workbook(path, data_only=True)
        if it['sheet'] not in wb.sheetnames:
            fail(f"{it['file']} 에 시트 '{it['sheet']}' 가 없음")
        ws = wb[it['sheet']]
        rows = parse_blocks(ws, region, group)
        if rows is None:
            rows = parse_daily4(ws, region, group, int(it.get('year', 2026)))
        if not rows:
            fail(f"{it['file']} / {it['sheet']}: 읽은 줄이 없음 (표 모양을 모름)")
        rows = [(x[0], x[1], name_fix.get(x[2], x[2])) + x[3:] for x in rows if x[1] in region[group]]
        days = sorted({x[0] for x in rows})
        print(f"   [배송 스케줄] {region['name']} {'주간' if group == 'day' else '야간'} ← {it['file']} / {it['sheet']}: {days[0]} ~ {days[-1]} {len(rows)}줄")
        for x in rows:
            k = (x[0], x[1], x[6])                         # 날짜·캠프·아이디
            if k in merged and merged[k][:6] != x[:6]:
                notes.append(f'{x[0]} {x[2]}: {merged[k][3:6]} → {x[3:6]} (뒤 표가 이김)')
            merged[k] = x
            src_t[k] = os.path.getmtime(path)
    for n in notes[:10]:
        print('     (참고) 표끼리 다름', n)
    when = {}
    for k, t in src_t.items():
        when[k[0]] = max(when.get(k[0], 0), t)
    return list(merged.values()), when


def disambiguate(rows):
    """같은 날 한 이름에 아이디가 둘 이상이면(동명이인) 이름 뒤에 표시를 붙이고, 아이디를 뺀 6칸 줄로 돌려줌.
    표시 = 주 라우트(서로 다르면) — 김상우(607A)·김상우(608A) / 주 라우트가 같으면 번호 — 정인호(1)·정인호(2).
    주마다 아이디가 바뀐 한 사람(같은 날 겹치지 않음)은 그대로 한 사람으로 둠."""
    day_ids = collections.defaultdict(set)
    for x in rows:
        day_ids[(x[0], x[2])].add(x[6])
    names = sorted({n for (d, n), s in day_ids.items() if len(s) > 1})
    rename = {}
    for name in names:
        pids = sorted({x[6] for x in rows if x[2] == name})
        top = {}
        for pid in pids:
            c = collections.Counter(','.join(x[5]) for x in rows if x[6] == pid and x[3] == '출근' and x[5])
            top[pid] = c.most_common(1)[0][0].split(',')[0] if c else ''
        distinct = len(set(top.values())) == len(pids) and all(top.values())
        for k, pid in enumerate(pids, 1):
            rename[(name, pid)] = f'{name}({top[pid]})' if distinct else f'{name}({k})'
        print(f'     (동명이인) {name} → {", ".join(rename[(name, p)] for p in pids)}')
    return [(x[0], x[1], rename.get((x[2], x[6]), x[2]), x[3], x[4], x[5]) for x in rows]


def build(rows, camps, wave, full_rot, title, when=None):
    if not rows:
        return None
    rows = disambiguate(rows)
    cd.CAMPS, cd.WAVE, cd.FULL_ROT = camps, wave, full_rot
    by_date = collections.defaultdict(list)
    for x in rows:
        by_date[x[0]].append(x)
    dates = sorted(by_date)

    def sunday(d):
        x = datetime.date.fromisoformat(d)
        return (x - datetime.timedelta(days=(x.weekday() + 1) % 7)).isoformat()
    week_dates = collections.defaultdict(list)
    for d in dates:
        week_dates[sunday(d)].append(d)
    weeks = [cd.build_week(wd, by_date) for _, wd in sorted(week_dates.items())]
    show = lambda c: re.sub(r'^[MV]_', '', c)              # 화면 캠프 이름: M_통영1 → 통영1 (어드민 다운로드 표기, 이름 칸에 들어가게)
    for w in weeks:
        for c in w['camps']:
            c['name'] = show(c['name'])
            for pp in c['people']:
                for v in pp['days']:
                    if isinstance(v, dict) and 'c' in v:
                        v['c'] = show(v['c'])
    for w in weeks:                                         # 마지막 업데이트(주마다) — 화면에 '마지막 업데이트 10/21(수) 18:38'
        ts = [when[d] for d in w['dates'] if when and d in when]
        if ts:
            w['updated'] = datetime.datetime.fromtimestamp(max(ts)).strftime('%Y-%m-%d %H:%M')
    return {'title': title, 'dates': dates, 'weeks': weeks}


def main():
    plan_cfg, name_fix = [], {}
    pj = os.path.join(SRC, 'plan.json')
    if os.path.exists(pj):
        cfg = json.load(io.open(pj, encoding='utf-8'))
        plan_cfg, name_fix = (cfg, {}) if isinstance(cfg, list) else (cfg.get('files', []), cfg.get('name_fix', {}))
    if name_fix:
        print('[이름 보정]', name_fix)
    out = []
    for rg in REGIONS:
        entry = {'key': rg['key'], 'name': rg['name']}
        for group in ('day', 'night'):
            wave, full_rot = WAVES[group]
            lab = '주간' if group == 'day' else '야간'
            camps = rg[group]
            arows, awhen = admin_rows(camps, wave)
            adm = build(arows, camps, wave, full_rot, f"{rg['name']} {lab} 전체스케쥴(어드민)", awhen)
            items = [x for x in plan_cfg if x['region'] == rg['name'] and x['group'] == group]
            pln = None
            if items:
                prows, pwhen = plan_rows(rg, group, items, name_fix)
                pln = build(prows, camps, wave, full_rot, f"{rg['name']} {lab} 배송 스케줄", pwhen)
            entry[group] = {'plan': pln, 'admin': adm}
            for k, v in (('배송 스케줄', pln), ('어드민', adm)):
                if v:
                    print(f"[{rg['name']} {lab} {k}] {v['dates'][0]} ~ {v['dates'][-1]} | " + ' / '.join(
                        f"{w['dates'][0][5:]}주 " + ' · '.join(f"{c['name']} {len(c['people'])}명" for c in w['camps']) for w in v['weeks']))
                else:
                    print(f"[{rg['name']} {lab} {k}] 자료 없음")
        out.append(entry)
    data = {'generated': datetime.date.today().isoformat(), 'regions': out}
    with io.open(OUT, 'w', encoding='utf-8', newline='\n') as f:
        f.write('// 자동 생성 파일 — 직접 수정하지 말고 tools/convert_region.py 로 다시 만드세요 (이름·라우트·휴무만, 아이디 없음)\n')
        f.write('window.REGION_DATA = ')
        json.dump(data, f, ensure_ascii=False, separators=(',', ':'))
        f.write(';\n')
    print(f'[저장] {OUT} ({os.path.getsize(OUT):,} bytes)')


if __name__ == '__main__':
    main()
