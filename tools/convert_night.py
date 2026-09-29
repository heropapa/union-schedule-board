# -*- coding: utf-8 -*-
"""야간 근무표(월별 xlsx) → night-data.js 변환 + 검증

사용법:
  python tools/convert_night.py <월별.xlsx> [<월별2.xlsx> ...] [<어드민양식.zip 또는 폴더>] [--force]

  예) python tools/convert_night.py 야간근무표_월별_2026.10~2027.01.xlsx 야간근무표_어드민양식_15주_주별.zip

■ 매달 절차
  1. 새 월별 엑셀을 받는다. 출력(night-data.js)은 통째로 새로 만들어지므로
     **이번 주를 포함한 보여줄 기간 전체**가 입력에 들어 있어야 한다.
     새 파일에 이번 달이 없으면 지난 파일도 같이 넣는다(여러 개 가능, 겹치는 날은 값이 같아야 함).
  2. 어드민 양식 zip(또는 폴더)이 있으면 같이 넣는다 → 한 칸이라도 다르면 중단.
  3. 성공하면 night-data.js 가 갱신된다. 기존 파일과 비교해 '오늘이 빠짐'·'앞으로의 날짜가 줄어듦'이면
     중단한다(정말 의도한 것이면 --force).
  4. 배포:  git add night-data.js && git commit -m "야간 근무표 갱신" && git push
     → Netlify 자동 배포 (1분 내).  확인: https://gleaming-mermaid-cc23e0.netlify.app/?g=night

■ 엑셀 규칙 (어기면 명확한 오류로 중단 — 조용히 틀린 데이터가 나가지 않게)
  - 시트명 'YYYY년 M월'. 그 외 시트(메모 등)는 무시.
  - '이름' 헤더 행 / 바로 아래 행에 날짜(숫자) / 헤더 행의 요일(일~토)은 실제 요일과 같아야 함.
  - 조 헤더 행: A열=조 이름, B열 비움, 날짜칸 비움.   인원 행: A열=이름, B열=담당(708A / 백업1 …).
  - 날짜칸 값: '휴' 또는 라우트(708A, 709AD, 520BCD …). 공백·소문자는 자동 정리,
    '휴무/연차/휴일/OFF' 는 '휴'로 바꾸고 안내를 출력. 그 외 값·중간 빈칸은 오류.
  - 표 아래 메모 행 금지(조로 오인됨). 동명이인은 이름을 구분(예: 김정철A).

■ 오류별 대처
  - '형식 오류/빈칸'      → 엑셀 해당 칸 수정
  - '요일 불일치'        → 시트명(연·월)이 틀렸거나 다른 달 탭을 복사한 것
  - '날짜 누락'          → 중간 달 시트가 빠짐
  - '어드민 불일치'       → 월별과 어드민 중 틀린 쪽을 확인(사장님께 문의)
  - '같은 이름이 두 번'   → 동명이인이면 이름 구분, 아니면 중복 행 삭제

필요: pip install openpyxl
"""
import sys, os, re, io, json, zipfile, datetime
import openpyxl

sys.stdout.reconfigure(encoding='utf-8')
ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
OUT = os.environ.get('NIGHT_OUT') or os.path.join(ROOT, 'night-data.js')
WD = '월화수목금토일'                        # date.weekday(): 월=0
OFF_ALIASES = {'휴무', '휴일', '연차', 'OFF', '오프'}
ROUTE_RE = re.compile(r'\d{3}[A-Z]+')


def fail(msg, details=()):
    for d in list(details)[:40]:
        print('   ✗', d)
    if len(details) > 40:
        print(f'   … 외 {len(details) - 40}건')
    raise SystemExit(f'[중단] {msg} — night-data.js를 쓰지 않았습니다.')


def norm_value(v, notes, where):
    s = re.sub(r'\s+', '', str(v)).upper()
    if s in OFF_ALIASES:
        notes.append(f"{where}: '{v}' → '휴' 로 변환")
        return '휴'
    if s in ('휴', ''):
        return s or None
    return s


def read_sheet(ws, y, mo):
    """시트 하나 → (제목, 안내, [(조이름, [(이름, 담당, {날짜: 값})])])"""
    hdr = next((r for r in range(1, min(ws.max_row, 30) + 1)
                if str(ws.cell(r, 1).value or '').strip() == '이름'), None)
    if hdr is None:
        fail(f"'{ws.title}' 시트에서 '이름' 헤더를 찾지 못함")
    title = note = None
    for r in range(1, hdr):
        t = str(ws.cell(r, 1).value or '').strip()
        if not t:
            continue
        if title is None and '—' in t:
            title = t.split('—')[0].strip()
        elif note is None and '·' in t:
            note = ' · '.join(x.strip() for x in t.split('·')
                              if x.strip() and not any(k in x for k in ('빨강', '파랑', '노랑')))
    daycols, errs = [], []
    for c in range(3, ws.max_column + 1):
        v = ws.cell(hdr + 1, c).value
        if isinstance(v, (int, float)) or (isinstance(v, str) and v.strip().isdigit()):
            try:
                d = datetime.date(y, mo, int(v))
            except ValueError:
                fail(f"'{ws.title}' 시트에 {int(v)}일이 있음 — 시트명(연·월)을 확인")
            wd = str(ws.cell(hdr, c).value or '').strip()
            if wd and wd[:1] in WD and wd[:1] != WD[d.weekday()]:
                errs.append(f"'{ws.title}' {d.day}일: 헤더 요일 '{wd}' ≠ 실제 '{WD[d.weekday()]}'")
            daycols.append((c, d.isoformat()))
    if errs:
        fail('요일 불일치 (시트명 연·월이 틀렸거나 다른 달 탭을 복사함)', errs)
    if not daycols:
        fail(f"'{ws.title}' 시트에서 날짜(숫자) 행을 찾지 못함")

    groups, cur, seen = [], None, set()
    for r in range(hdr + 2, ws.max_row + 1):
        a, b = ws.cell(r, 1).value, ws.cell(r, 2).value
        a = str(a).strip() if a is not None else ''
        b = str(b).strip() if b is not None else ''
        vals = {d: ws.cell(r, c).value for c, d in daycols}
        has_vals = any(v not in (None, '') and str(v).strip() for v in vals.values())
        if not a:
            if b or has_vals:
                fail(f"'{ws.title}' {r}행: 이름(A열)이 비었는데 값이 있음")
            continue
        if not b:
            if has_vals:
                fail(f"'{ws.title}' {r}행 '{a}': 담당(B열)이 비었는데 날짜칸에 값이 있음")
            if '조' not in a:
                fail(f"'{ws.title}' {r}행 '{a}': 조 헤더로 보기 어려움 (조 이름엔 '조'가 들어가야 함. 메모 행이면 표 밖으로)")
            cur = (a, [])                     # 조 헤더
            groups.append(cur)
            continue
        if cur is None:
            fail(f"'{ws.title}' {r}행 '{a}': 조 헤더(예: 'A조 (화전)') 없이 인원 행이 나옴")
        if a in seen:
            fail(f"'{ws.title}' 시트에 같은 이름 '{a}'이(가) 두 번 — 동명이인이면 이름을 구분(예: {a}A)")
        seen.add(a)
        cur[1].append((a, b, vals))
    return title, note, groups


def parse_monthly(paths):
    sheets = []                               # (y, mo, 파일, 시트)
    for path in paths:
        wb = openpyxl.load_workbook(path, data_only=True)
        for ws in wb.worksheets:
            m = re.fullmatch(r'\s*(\d{4})년\s*(\d{1,2})월\s*', ws.title)
            if m:
                sheets.append((int(m.group(1)), int(m.group(2)), os.path.basename(path), ws))
            else:
                print(f"   (무시) '{os.path.basename(path)}' 의 '{ws.title}' 시트 — 'YYYY년 M월' 형식 아님")
    if not sheets:
        fail("'YYYY년 M월' 형식 시트가 없음 (월별 근무표 파일이 맞는지 확인)")
    sheets.sort(key=lambda s: (s[0], s[1]))

    title = note = None
    cells, roles, team_of, notes = {}, {}, {}, []
    order = []                                # 시트별 [(조, [이름...])]
    for y, mo, fname, ws in sheets:
        t, n, groups = read_sheet(ws, y, mo)
        title = title or t
        note = note or n
        order.append([(g, [p[0] for p in ps]) for g, ps in groups])
        for g, ps in groups:
            for name, role, vals in ps:
                if name in team_of and team_of[name] != g:
                    notes.append(f"{name}: 조 변경 '{team_of[name]}' → '{g}' ({ws.title}) — 최신 달 기준으로 표시")
                team_of[name] = g
                for d, raw in vals.items():
                    if raw is None or str(raw).strip() == '':
                        continue
                    v = norm_value(raw, notes, f'{name} {d}')
                    if v is None:
                        continue
                    if (name, d) in cells and cells[(name, d)] != v:
                        fail(f'같은 날 값이 파일/시트마다 다름: {name} {d} {cells[(name, d)]} vs {v} ({fname} {ws.title})')
                    cells[(name, d)] = v
                    roles[(name, d)] = role

    dates = sorted({d for _, d in cells})
    if not dates:
        fail('근무 값이 하나도 없음')
    d0 = datetime.date.fromisoformat(dates[0])
    for i, d in enumerate(dates):
        exp = (d0 + datetime.timedelta(days=i)).isoformat()
        if exp != d:
            fail(f'날짜 누락: {exp} 부터 비어 있음 (중간 달 시트가 빠졌는지 확인)')

    # 조/인원 순서: 최신 달 구성 기준, 이전 달에만 있는 사람(퇴사 등)은 그 조 끝에 덧붙임
    teams = []
    for snap in reversed(order):
        for g, names in snap:
            t = next((x for x in teams if x['name'] == g), None)
            if t is None:
                t = {'name': g, 'people': []}
                teams.append(t)
            for nm in names:
                if team_of[nm] == g and not any(p['name'] == nm for p in t['people']):
                    t['people'].append({'name': nm})
    kept = []
    for t in teams:
        if t['people']:
            kept.append(t)
        else:
            notes.append(f"인원 없는 조 '{t['name']}' 제외 (이전 달에 쓰던 조 이름?)")
    teams = kept
    listed = {p['name'] for t in teams for p in t['people']}
    if listed != set(team_of):
        fail('조 배치 오류', sorted(set(team_of) - listed))

    errs = []
    for t in teams:
        for p in t['people']:
            nm = p['name']
            p['days'] = [cells.get((nm, d)) for d in dates]
            rl = [roles.get((nm, d)) for d in dates]
            idx = [i for i, v in enumerate(p['days']) if v is not None]
            if not idx:
                errs.append(f"근무 값이 하나도 없는 사람: '{nm}'")
                continue
            for i in range(idx[0], idx[-1] + 1):
                v = p['days'][i]
                if v is None:
                    errs.append(f'빈칸: {nm} {dates[i]}')
                elif v != '휴' and not ROUTE_RE.fullmatch(v):
                    errs.append(f"형식 오류: {nm} {dates[i]} '{v}' ('휴' 또는 708A 같은 라우트)")
            if idx[0] > 0 or idx[-1] < len(dates) - 1:
                notes.append(f'{nm}: {dates[idx[0]]} ~ {dates[idx[-1]]} 만 근무표에 있음 (입사/퇴사?)')
            used = [r for r in rl if r is not None]
            p['role'] = used[-1]                                  # 최신 담당
            if len(set(used)) > 1:                                # 달마다 담당이 바뀐 경우 날짜별로 보관
                p['roles'] = rl
                notes.append(f"{nm}: 담당 변경 {' → '.join(dict.fromkeys(used))}")
    if errs:
        fail(f'월별 검증 실패 {len(errs)}건', errs)
    for t in teams:                           # 출력 키 순서 정리
        t['people'] = [{k: p[k] for k in ('name', 'role', 'roles', 'days') if k in p} for p in t['people']]
    return title or '야간 근무표', note or '', dates, teams, notes


def expand(code):
    """'709AD' → {'709A','709D'}"""
    m = re.fullmatch(r'(\d+)([A-Z]+)', code)
    return {m.group(1) + ch for ch in m.group(2)} if m else {code}


def load_admin(src):
    skip = lambda n: '__MACOSX' in n or os.path.basename(n).startswith(('._', '~$'))
    files = []
    if src.lower().endswith('.zip'):
        z = zipfile.ZipFile(src)
        for i in z.infolist():
            if i.filename.lower().endswith('.xlsx') and not skip(i.filename):
                files.append(io.BytesIO(z.read(i)))
    else:
        files = [os.path.join(src, f) for f in sorted(os.listdir(src))
                 if f.lower().endswith('.xlsx') and not skip(f)]
    if not files:
        fail(f'어드민 양식 xlsx를 찾지 못함: {src}')
    rows = []
    for f in files:
        ws = openpyxl.load_workbook(f, data_only=True).worksheets[0]
        hdr = [c.value for c in ws[1]]
        ix = {h: i for i, h in enumerate(hdr)}
        for need in ('업무일', '이름', '업무상태', '업무라우트'):
            if need not in ix:
                fail(f"어드민 양식에 '{need}' 열이 없음")
        for r in ws.iter_rows(min_row=2, values_only=True):
            if r[ix['업무일']] is None:
                continue
            d = r[ix['업무일']]
            d = d.date().isoformat() if hasattr(d, 'date') else str(d)[:10]
            rows.append((d, str(r[ix['이름']]).strip(), str(r[ix['업무상태']]).strip(),
                         re.sub(r'\s+', '', str(r[ix['업무라우트']] or '')).upper()))
    return rows


def cross_check(dates, teams, admin_rows):
    grid = {(p['name'], d): v for t in teams for p in t['people'] for d, v in zip(dates, p['days'])}
    bad, seen = [], set()
    for d, name, st, routes in admin_rows:
        seen.add((name, d))
        v = grid.get((name, d))
        if v is None:
            bad.append(f'{d} {name}: 월별에 없음 (어드민 {st} {routes})')
        elif st == '휴무':
            if v != '휴':
                bad.append(f'{d} {name}: 어드민 휴무 vs 월별 {v}')
        elif v == '휴' or expand(v) != {x for x in routes.split(',') if x}:
            bad.append(f'{d} {name}: 어드민 {st} {routes} vs 월별 {v}')
    adates = sorted({d for d, *_ in admin_rows})
    for (name, d), v in grid.items():
        if v is not None and d in adates and (name, d) not in seen:
            bad.append(f'{d} {name}: 어드민에 없음 (월별 {v})')
    return adates, bad


def load_existing():
    if not os.path.exists(OUT):
        return None
    s = io.open(OUT, encoding='utf-8').read()
    try:
        return json.loads(s[s.index('=') + 1:].strip().rstrip(';'))
    except Exception:
        return None


def main():
    args = [a for a in sys.argv[1:] if a != '--force']
    force = '--force' in sys.argv[1:]
    monthly = [a for a in args if a.lower().endswith('.xlsx')]
    admin = [a for a in args if a.lower().endswith('.zip') or os.path.isdir(a)]
    other = [a for a in args if a not in monthly and a not in admin]
    if not monthly or other or len(admin) > 1:
        raise SystemExit(__doc__ + (f'\n알 수 없는 인자: {other}' if other else ''))

    title, note, dates, teams, notes = parse_monthly(monthly)
    n_people = sum(len(t['people']) for t in teams)
    print(f'[월별] {title} | {dates[0]} ~ {dates[-1]} ({len(dates)}일) | {len(teams)}개 조 {n_people}명')
    for n in notes:
        print('   (참고)', n)

    if admin:
        adates, bad = cross_check(dates, teams, load_admin(admin[0]))
        print(f'[어드민 대조] {adates[0]} ~ {adates[-1]} ({len(adates)}일) → 불일치 {len(bad)}건')
        if bad:
            fail('월별과 어드민이 다름', bad)
        uncovered = [d for d in dates if d not in set(adates)]
        if uncovered:
            print(f'   (참고) 어드민에 없는 {len(uncovered)}일({uncovered[0]} ~ {uncovered[-1]} 등)은 월별 규칙 검사만 거침')
    else:
        print('[어드민 대조] 생략 (zip을 같이 넣으면 한 칸씩 대조합니다)')

    old = load_existing()
    if old and old.get('dates'):
        today = os.environ.get('NIGHT_TODAY') or datetime.date.today().isoformat()   # NIGHT_TODAY: 테스트용
        o0, o1 = old['dates'][0], old['dates'][-1]
        stop = []
        if o0 <= today <= o1 and not (dates[0] <= today <= dates[-1]):
            stop.append(f'오늘({today})이 기존 근무표엔 있는데 새 근무표({dates[0]}~{dates[-1]})엔 없음 → 이번 주가 사라짐')
        if dates[-1] < o1:
            stop.append(f'끝나는 날이 앞당겨짐: {o1} → {dates[-1]} (앞으로의 근무표가 줄어듦)')
        og = {(p['name'], d): v for t in old['teams'] for p in t['people'] for d, v in zip(old['dates'], p['days'])}
        ng = {(p['name'], d): v for t in teams for p in t['people'] for d, v in zip(dates, p['days'])}
        changed = [k for k in ng if k in og and og[k] != ng[k] and og[k] is not None]
        on, nn = {k[0] for k in og}, {k[0] for k in ng}
        print(f'[기존 대비] 기간 {o0}~{o1} → {dates[0]}~{dates[-1]} | 바뀐 칸 {len(changed)}'
              + (f' | 추가 {sorted(nn - on)}' if nn - on else '') + (f' | 빠짐 {sorted(on - nn)}' if on - nn else ''))
        for k in changed[:10]:
            print(f'   (변경) {k[0]} {k[1]}: {og[k]} → {ng[k]}')
        if stop and not force:
            fail('기존 근무표보다 보여줄 기간이 줄어듦 (정말 의도했다면 --force)', stop)

    data = {'title': title, 'note': note, 'source': ', '.join(os.path.basename(m) for m in monthly),
            'generated': datetime.date.today().isoformat(), 'dates': dates, 'teams': teams}
    with io.open(OUT, 'w', encoding='utf-8', newline='\n') as f:
        f.write('// 자동 생성 파일 — 직접 수정하지 말고 tools/convert_night.py 로 다시 만드세요\n')
        f.write('window.NIGHT_DATA = ')
        json.dump(data, f, ensure_ascii=False, separators=(',', ':'))
        f.write(';\n')
    print(f'[저장] {OUT} ({os.path.getsize(OUT):,} bytes)')


if __name__ == '__main__':
    main()
