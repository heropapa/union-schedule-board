# -*- coding: utf-8 -*-
"""야간 근무표(기준 데이터) → 쿠팡 어드민 업로드 엑셀 (부산2·부산3 WAVE1)

사용법:
  python tools/make_night_admin.py <PDD 시작일 YYYY-MM-DD> [<PDD 끝일>] [--out <폴더>]
  예) python tools/make_night_admin.py 2026-10-11              → 10/11~10/17 (7일)
      python tools/make_night_admin.py 2026-10-12 2026-10-17   → 그 기간만 (마감 지난 날 빼고 올릴 때)

- 기준 데이터: source/수정이력.md 의 재생성 명령과 같은 입력(월별 수정본 + 15주 어드민 zip + 9/27~10/3 어드민 수정본)으로
  convert_night.py 를 공개 범위 제한 없이 돌린 결과 (임시 파일에만, night-data.js 는 안 건드림).
- 아이디·벤더명·사업자등록번호: source/ 의 쿠팡 내보내기에서 읽음 → 이 파일·저장소엔 안 들어감 (repo 공개).
- 규칙 (기존 15주 어드민 zip 1,050줄을 한 줄도 안 틀리고 재현하는지 매번 먼저 확인 — 다르면 중단):
  한 사람·하루 1줄 / 정렬 = 캠프(부산2→부산3) → 출근 먼저 → 근무표 조 순서 /
  출근 캠프 = 라우트 번호(708·709→부산2, 그 외→부산3), 회전 'D1,D2,F3'(일부 회전이면 그대로) /
  휴무 = 회전·라우트 빈칸, 캠프는 조 기준(A조→부산2, B조→부산3 — 어드민이 받아 줌, 2026-10-01 확인).
- 서식: zip 의 업로드 파일을 틀로 열어 데이터 행만 바꿔 씀 (다른 통합문서의 서식 복사 금지 — 엑셀이 못 여는 파일이 됨).
- 날짜 = PDD(업무일). 수정 마감은 업무 전날 오전 10시(W1).
"""
import sys, io, os, re, json, zipfile, datetime, subprocess, tempfile, copy, warnings
import openpyxl

warnings.filterwarnings('ignore')
sys.stdout.reconfigure(encoding='utf-8')
ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
S = os.path.join(ROOT, 'source')
MONTHLY = os.path.join(S, '야간근무표_월별_2026.10~2027.01_수정본.xlsx')
ZIP = os.path.join(S, '야간근무표_어드민양식_15주_주별_수정본.zip')
EXPORT = os.path.join(S, '어드민_2026-09-27~10-03_schedule-v2_수정본.xlsx')
CAMP_ORDER = {'부산2': 0, '부산3': 1}


def main():
    args = sys.argv[1:]
    out_dir = os.path.join(S, 'night_admin_upload')
    if '--out' in args:
        i = args.index('--out'); out_dir = args[i + 1]; del args[i:i + 2]
    if not args:
        raise SystemExit(__doc__)
    d0 = datetime.date.fromisoformat(args[0])
    d1 = datetime.date.fromisoformat(args[1]) if len(args) > 1 else d0 + datetime.timedelta(days=6)
    dates = [(d0 + datetime.timedelta(days=i)).isoformat() for i in range((d1 - d0).days + 1)]

    # 기준 데이터 전체 (공개 범위 제한 없이, 임시 파일)
    tmp = os.path.join(tempfile.gettempdir(), 'night_full_for_admin.js')
    env = dict(os.environ, NIGHT_OUT=tmp, PYTHONUTF8='1')
    r = subprocess.run([sys.executable, os.path.join(ROOT, 'tools', 'convert_night.py'), MONTHLY, ZIP, EXPORT],
                       cwd=ROOT, env=env, capture_output=True, text=True, encoding='utf-8')
    if r.returncode:
        raise SystemExit('[중단] 기준 데이터 만들기 실패\n' + r.stdout + r.stderr)
    s = io.open(tmp, encoding='utf-8').read(); ND = json.loads(s[s.index('=') + 1:].strip().rstrip(';'))
    order = [(t['name'], p['name']) for t in ND['teams'] for p in t['people']]
    val = {(p['name'], d): v for t in ND['teams'] for p in t['people'] for d, v in zip(ND['dates'], p['days'])}
    team_of = {n: t for t, n in order}
    pos = {n: i for i, (t, n) in enumerate(order)}
    miss = [d for d in dates if not any((n, d) in val for _, n in order)]
    if miss:
        raise SystemExit(f'[중단] 기준 데이터에 없는 날짜: {miss[0]} ~ {miss[-1]}')

    # 아이디·벤더·사업자번호
    ws = openpyxl.load_workbook(EXPORT, read_only=True, data_only=True).worksheets[0]
    rows = list(ws.iter_rows(values_only=True)); ix = {h: i for i, h in enumerate(rows[0])}
    ids = {}; vendor = biz = None
    for x in rows[1:]:
        if x[ix['이름']]:
            ids.setdefault(x[ix['이름']], set()).add(x[ix['아이디']]); vendor, biz = x[ix['벤더명']], x[ix['사업자등록번호']]
    if any(len(v) != 1 for v in ids.values()):
        raise SystemExit('[중단] 이름당 아이디가 하나가 아님')
    ids = {k: next(iter(v)) for k, v in ids.items()}
    nohave = sorted({n for _, n in order} - set(ids))
    if nohave:
        raise SystemExit(f'[중단] 아이디를 모르는 사람: {nohave} (최근 야간 어드민 내보내기를 source/ 에 넣어 주세요)')

    def rows_for(ds):
        out = []
        for d in ds:
            day = []
            for t, n in order:
                v = val.get((n, d))
                if v is None: continue
                if v == '휴':
                    day.append(('부산2' if team_of[n].startswith('A') else '부산3', 1, pos[n], d, n, '휴무', None, None))
                else:
                    m = re.fullmatch(r'(\d{3})([A-Z]+)(?:\((.+)\))?', v)
                    camp = '부산2' if m.group(1) in ('708', '709') else '부산3'
                    day.append((camp, 0, pos[n], d, n, '출근', m.group(3) or 'D1,D2,F3', ','.join(m.group(1) + c for c in m.group(2))))
            day.sort(key=lambda x: (CAMP_ORDER[x[0]], x[1], x[2]))
            out += day
        return out

    # 규칙 검증: 기존 15주 zip 재현
    z = zipfile.ZipFile(ZIP); bad = []; total = 0; tpl_name = None
    for nm in sorted(z.namelist()):
        if not nm.endswith('.xlsx') or os.path.basename(nm).startswith('._'): continue
        tpl_name = tpl_name or nm
        w = openpyxl.load_workbook(io.BytesIO(z.read(nm)), data_only=True).worksheets[0]
        rr = list(w.iter_rows(values_only=True)); hx = {h: i for i, h in enumerate(rr[0])}
        zr = [(x[hx['업무일']].date().isoformat(), x[hx['캠프명']], x[hx['이름']], x[hx['업무상태']], x[hx['회전']], x[hx['업무라우트']]) for x in rr[1:] if x[hx['이름']]]
        mine = [(d, c, n, st, rot, rt) for (c, _, _, d, n, st, rot, rt) in rows_for(sorted({x[0] for x in zr}))]
        total += len(zr)
        if mine != zr: bad.append(nm)
        bad += [f'{nm} 아이디 다름 {x[hx["이름"]]}' for x in rr[1:] if x[hx['이름']] and ids.get(x[hx['이름']]) != x[hx['아이디']]]
    if bad:
        raise SystemExit(f'[중단] 규칙이 기존 15주 양식과 다름: {bad[:5]}')
    print(f'[규칙 확인] 기존 15주 양식 {total}줄 그대로 재현')

    # 생성 (zip 업로드 파일을 틀로)
    tpl = openpyxl.load_workbook(io.BytesIO(z.read(tpl_name)))
    w = tpl.worksheets[0]
    st = [copy.copy(w.cell(2, c)._style) for c in range(1, 11)]
    w.delete_rows(2, w.max_row)
    new = rows_for(dates)
    for i, (c, _, _, d, n, stt, rot, rt) in enumerate(new, start=2):
        for j, v in enumerate([datetime.datetime.fromisoformat(d), vendor, biz, c, 'WAVE1', n, ids[n], stt, rot, rt], start=1):
            cell = w.cell(i, j, v); cell._style = copy.copy(st[j - 1])
        w.cell(i, 1).number_format = 'yyyy/MM/dd'
    os.makedirs(out_dir, exist_ok=True)
    out = os.path.join(out_dir, f'야간_어드민업로드_{dates[0][5:]}~{dates[-1][5:]}.xlsx')
    tpl.save(out)
    print(f'[저장] {out} | {len(new)}줄 ({dates[0]} ~ {dates[-1]}, {len(dates)}일)')


if __name__ == '__main__':
    main()
